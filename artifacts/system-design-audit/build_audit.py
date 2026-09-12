from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
import shutil
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path('/Users/miladsaki/Downloads/personal-wealth-operating-system')
REFERENCE = Path('/Users/miladsaki/.codex/plugins/cache/openai-curated-remote/openai-templates/0.1.1/skills/artifact-template-system-design/assets/reference.docx')
OUT = ROOT / 'artifacts/system-design-audit/financial-system-design-template-audit-fa.docx'
BEFORE = Path('/private/tmp/system-design-audit-20260912/reference-render/page-1.png')
LOGO = ROOT / 'public/icon-512.png'

INK = '101218'
NAVY = '14233B'
IRIS = '4F46E5'
IRIS_SOFT = 'EEF0FF'
SLATE = '526173'
MUTED = '6B7280'
LINE = 'D9DCE3'
PALE = 'F5F6F8'
GREEN = '138A61'
GREEN_SOFT = 'E8F6F0'
AMBER = 'A66B12'
AMBER_SOFT = 'FFF4DB'
RED = 'B13A3A'
RED_SOFT = 'FCECEC'
WHITE = 'FFFFFF'


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for margin, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{margin}'))
        if node is None:
            node = OxmlElement(f'w:{margin}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_cell_borders(cell, color=LINE, size='6'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'start', 'bottom', 'end', 'insideH', 'insideV'):
        tag = f'w:{edge}'
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), size)
        node.set(qn('w:color'), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn('w:keepNext'))
    if node is None:
        node = OxmlElement('w:keepNext')
        p_pr.append(node)
    node.set(qn('w:val'), '1' if value else '0')


def set_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn('w:bidi'))
    if bidi is None:
        bidi = OxmlElement('w:bidi')
        p_pr.append(bidi)
    bidi.set(qn('w:val'), '1')


def set_run_font(run, name='Arial', size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:cs'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(p, text, *, bold=False, color=INK, size=10.5):
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    return r


def set_para_spacing(p, before=0, after=6, line=1.15):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_body(doc, text, *, bold_lead=None, after=7, size=10.5, color=INK):
    p = doc.add_paragraph(style='Normal')
    set_rtl(p)
    set_para_spacing(p, after=after, line=1.24)
    if bold_lead and text.startswith(bold_lead):
        add_text(p, bold_lead, bold=True, size=size, color=color)
        add_text(p, text[len(bold_lead):], size=size, color=color)
    else:
        add_text(p, text, size=size, color=color)
    return p


def add_heading(doc, text, level=1, kicker=None):
    if kicker:
        p0 = doc.add_paragraph()
        set_rtl(p0)
        set_para_spacing(p0, before=5, after=2)
        add_text(p0, kicker, bold=True, color=IRIS, size=8.5)
        set_keep_with_next(p0)
    p = doc.add_paragraph(style=f'Heading {level}')
    set_rtl(p)
    set_para_spacing(p, before=7 if level == 1 else 5, after=6 if level == 1 else 4)
    size = 20 if level == 1 else 13
    add_text(p, text, bold=True, color=INK, size=size)
    set_keep_with_next(p)
    return p


def add_rule(doc, color=IRIS, width='1280'):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '18')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    borders.append(bottom)
    p_pr.append(borders)
    p.paragraph_format.space_after = Pt(8)
    return p


def set_table_layout(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)


def fill_cell(cell, text, *, bg=WHITE, color=INK, bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT, size=9.5):
    set_cell_shading(cell, bg)
    p = cell.paragraphs[0]
    p.clear()
    set_rtl(p, align)
    set_para_spacing(p, after=0, line=1.15)
    add_text(p, text, bold=bold, color=color, size=size)


def add_matrix(doc, headers, rows, widths, header_fill=NAVY, body_fills=(WHITE, PALE), font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_layout(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        fill_cell(hdr.cells[i], h, bg=header_fill, color=WHITE, bold=True, size=9.2)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        fill = body_fills[ridx % len(body_fills)]
        for i, text in enumerate(row):
            fill_cell(cells[i], text, bg=fill, color=INK, size=font_size)
    p = doc.add_paragraph()
    set_para_spacing(p, after=2)
    return table


def add_bullet(doc, text, *, color=INK, marker='•'):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    set_para_spacing(p, after=4, line=1.18)
    add_text(p, f'{marker}  ', bold=True, color=IRIS, size=10)
    add_text(p, text, color=color, size=10)
    return p


def add_page_break(doc):
    doc.add_page_break()


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=8.5, color=MUTED)


def add_alt_text_to_images(docx_path: Path):
    with tempfile.TemporaryDirectory() as td:
        td_path = Path(td)
        with ZipFile(docx_path, 'r') as zin:
            zin.extractall(td_path)
        document_xml = td_path / 'word/document.xml'
        from lxml import etree
        tree = etree.parse(str(document_xml))
        ns = {'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
        descriptions = [
            'نشان اپ مدیریت ثروت شخصی توازن',
            'جلد نسخه موجود قالب عمومی System Design',
        ]
        for idx, node in enumerate(tree.xpath('//wp:docPr', namespaces=ns)):
            node.set('descr', descriptions[idx] if idx < len(descriptions) else 'تصویر سند ممیزی طراحی سیستم مالی')
        tree.write(str(document_xml), xml_declaration=True, encoding='UTF-8', standalone=True)
        tmp_out = docx_path.with_suffix('.patched.docx')
        with ZipFile(tmp_out, 'w', ZIP_DEFLATED) as zout:
            for file in sorted(td_path.rglob('*')):
                if file.is_file():
                    zout.write(file, file.relative_to(td_path))
        tmp_out.replace(docx_path)


def configure_styles(doc):
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Arial'
    normal._element.rPr.rFonts.set(qn('w:cs'), 'Arial')
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    for style_name, size in [('Title', 30), ('Heading 1', 20), ('Heading 2', 13), ('Heading 3', 11)]:
        st = styles[style_name]
        st.font.name = 'Arial'
        st._element.rPr.rFonts.set(qn('w:cs'), 'Arial')
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(INK)


def clear_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def setup_headers_footers(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    for header in (sec.header, sec.first_page_header):
        for p in header.paragraphs:
            p.clear()
    for footer in (sec.footer, sec.first_page_footer):
        for p in footer.paragraphs:
            p.clear()
    p = sec.footer.paragraphs[0]
    set_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, 'ممیزی قالب طراحی سیستم مالی  |  صفحه ', color=MUTED, size=8.5)
    add_page_field(p)
    p2 = sec.first_page_footer.paragraphs[0]
    set_rtl(p2, WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p2, 'نسخه ۱  |  ۲۱ شهریور ۱۴۰۵', color=MUTED, size=8.5)


def cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if LOGO.exists():
        run = p.add_run()
        run.add_picture(str(LOGO), width=Inches(0.62))
    p.paragraph_format.space_after = Pt(64)

    eyebrow = doc.add_paragraph()
    set_rtl(eyebrow)
    set_para_spacing(eyebrow, after=8)
    add_text(eyebrow, 'ممیزی طراحی و معماری سند', bold=True, color=IRIS, size=10)

    title = doc.add_paragraph(style='Title')
    set_rtl(title)
    set_para_spacing(title, after=10, line=1.05)
    add_text(title, 'قالب طراحی سیستم برای اپ مالی', bold=True, color=INK, size=30)

    sub = doc.add_paragraph()
    set_rtl(sub)
    set_para_spacing(sub, after=34, line=1.25)
    add_text(sub, 'ارزیابی نسخه موجود و پیشنهاد نسخه حرفه‌ای برای توازن', color=SLATE, size=13)

    add_rule(doc, color=IRIS)

    summary = doc.add_paragraph()
    set_rtl(summary)
    set_para_spacing(summary, after=26, line=1.35)
    add_text(summary, 'جمع‌بندی  ', bold=True, color=INK, size=11)
    add_text(summary, 'قالب فعلی از نظر ساختار مهندسی قابل اتکاست، اما برای یک اپ مالی فارسی به بازطراحی هویت، معماری اطلاعات و کنترل‌های مالی نیاز دارد. پیشنهاد این گزارش، حفظ اسکلت RFC و افزودن لایه‌های اعتماد مالی، حسابرسی‌پذیری و برند محصول است.', color=SLATE, size=11)

    meta = doc.add_table(rows=4, cols=2)
    set_table_layout(meta, [2.1, 4.8])
    set_repeat_table_header(meta.rows[0])
    fill_cell(meta.rows[0].cells[0], 'مشخصه', bg=NAVY, color=WHITE, bold=True, size=9.2)
    fill_cell(meta.rows[0].cells[1], 'مقدار', bg=NAVY, color=WHITE, bold=True, size=9.2)
    data = [
        ('وضعیت', 'پیشنهاد آماده تصمیم'),
        ('دامنه', 'قالب System Design و تناسب آن با سامانه مدیریت ثروت شخصی'),
        ('مبنای ارزیابی', 'فایل مرجع، رابط زنده محلی، اسناد طراحی و معماری مخزن'),
    ]
    for i, (label, value) in enumerate(data, start=1):
        fill_cell(meta.rows[i].cells[0], label, bg=IRIS_SOFT, bold=True, color=IRIS, size=9.5)
        fill_cell(meta.rows[i].cells[1], value, bg=WHITE, color=INK, size=9.5)


def executive_audit(doc):
    add_heading(doc, 'نتیجه ممیزی', kicker='خلاصه مدیریتی')
    add_body(doc, 'قالب موجود ۶۸ از ۱۰۰ امتیاز می‌گیرد. ساختار پوشش موضوعات فنی خوب است و الگوی جدول‌ها منسجم است، اما کیفیت برند، دسترس‌پذیری، معماری اطلاعات مالی و قابلیت استفاده در تصمیم‌گیری اجرایی فاصله محسوسی با یک سند ممتاز دارد.')

    scores = [
        ('ساختار مهندسی', '۸۴', 'قوی', 'پوشش معماری، قرارداد داده، خطا و عملیات'),
        ('هویت و برند', '۴۲', 'ضعیف', 'قالب عمومی و بدون پیوند با زبان بصری محصول'),
        ('تناسب مالی', '۵۱', 'نیازمند بازطراحی', 'فاقد کنترل‌های دفترکل، ارزش‌گذاری و تطبیق'),
        ('خوانایی و اسکن', '۷۶', 'خوب', 'جدول‌های منظم؛ جلد کم‌اطلاعات و برخی صفحات متراکم'),
        ('دسترس پذیری', '۵۸', 'ریسک متوسط', 'تصویر بدون متن جایگزین و پرش سطح تیتر'),
        ('آمادگی اجرا', '۶۹', 'متوسط', 'چک‌لیست عملیاتی دارد؛ شواهد پذیرش مالی کافی نیست'),
    ]
    add_matrix(doc, ['محور', 'امتیاز', 'وضعیت', 'شاهد اصلی'], scores, [1.55, 0.75, 1.35, 3.45], font_size=8.8)

    add_heading(doc, 'تصمیم پیشنهادی', level=2)
    add_body(doc, 'اسکلت ۱۲ بخشی حفظ شود، اما خروجی به یک قالب مالی تصمیم‌محور تبدیل شود. نسخه جدید باید در همان دو صفحه اول پاسخ دهد چه تصمیمی مطرح است، چه عددی منبع حقیقت است، چه کنترل‌هایی از ثبت دوباره یا دسترسی غیرمجاز جلوگیری می‌کنند و چه مدرکی برای عرضه لازم است.')
    add_bullet(doc, 'هویت بصری سند با محصول یکپارچه شود و نام نهایی برند پیش از انتشار تثبیت شود.')
    add_bullet(doc, 'بخش مستقل برای ثابت‌های مالی، مسیر حسابرسی، منشأ نرخ و ارزش‌گذاری و تطبیق دفترکل اضافه شود.')
    add_bullet(doc, 'ماتریس ریسک، مالک کنترل و شواهد پذیرش جایگزین متن‌های placeholder عمومی شود.')


def before_after(doc):
    add_heading(doc, 'مقایسه قبل و بعد', kicker='طرح پیشنهادی')
    rows = [
        ('جلد', 'عنوان عمومی و فضای خالی زیاد', 'عنوان تصمیم، وضعیت، دامنه، مبنای ارزیابی و هویت محصول در یک نگاه'),
        ('رنگ', 'سرمه‌ای و آبی سازمانی عمومی', 'جوهر تیره و Iris برند؛ سبز و قرمز فقط برای معنای مالی'),
        ('تایپوگرافی', 'Helvetica انگلیسی و قالب چپ‌چین', 'فونت فارسی خوانا، راست‌چین واقعی و اعداد جدولی'),
        ('ناوبری', '۱۲ بخش بدون نقشه تصمیم', 'خلاصه مدیریتی، امتیاز ریسک، سپس شواهد و تصمیم'),
        ('معماری', 'نمودار خطی Client تا Downstream', 'مسیر مستقل خواندن، نوشتن دفترکل، داده بازار و audit trail'),
        ('قرارداد داده', 'فهرست عمومی field و type', 'مبلغ، ارز، زمان مؤثر، نرخ منبع، مالک، idempotency key و نسخه'),
        ('امنیت', 'چک‌لیست عمومی Auth و Privacy', 'RLS مالک‌محور، مرز tenant، حفاظت کلید، session و دسترسی عملیاتی'),
        ('عملیات', 'SLO عمومی latency و success rate', 'تراز دفترکل، پوشش قیمت، خطای نرخ، RPO و RTO و آزمون بازیابی'),
        ('دسترس پذیری', 'تصویر بدون alt و پرش Heading 1 به 3', 'alt text، سلسله‌مراتب Heading 1 و 2 و جدول‌های قابل خواندن'),
    ]
    add_matrix(doc, ['مولفه', 'قبل', 'بعد پیشنهادی'], rows, [1.2, 2.55, 3.35], font_size=8.7)

    add_heading(doc, 'اثر مورد انتظار', level=2)
    add_body(doc, 'نسخه پیشنهادی زمان فهم سند را کاهش می‌دهد، چون تصمیم، ریسک و شواهد را جلوتر از جزئیات قرار می‌دهد. همچنین بررسی‌کننده مالی می‌تواند بدون جست‌وجو در متن تشخیص دهد عدد مرجع از کجا آمده، کدام عملیات برگشت‌پذیر است و چه کنترل‌هایی پیش از عرضه باید عبور کنند.')


def visual_evidence(doc):
    add_heading(doc, 'شواهد بصری نسخه موجود', kicker='قبل')
    add_body(doc, 'جلد فعلی تمیز است، اما بیش از نیمی از صفحه را بدون هدف رها می‌کند. هیچ نشانه‌ای از محصول مالی، مخاطب تصمیم، سطح محرمانگی یا وضعیت کنترل‌ها دیده نمی‌شود.')
    if BEFORE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(str(BEFORE), width=Inches(4.55))
        set_para_spacing(p, after=5)
        cap = doc.add_paragraph()
        set_rtl(cap, WD_ALIGN_PARAGRAPH.CENTER)
        set_para_spacing(cap, after=10)
        add_text(cap, 'نمونه جلد قالب مرجع', color=MUTED, size=8.5)

    add_heading(doc, 'جهت بصری نسخه پیشنهادی', level=2, kicker='بعد')
    add_body(doc, 'زبان پیشنهادی از خود محصول می‌آید: زمینه روشن یا جوهری، Iris برای برند، خطوط کم، تراکم کنترل‌شده و رنگ‌های مثبت و منفی فقط برای وضعیت مالی. صفحه اول همین گزارش نمونه اجرایی این جهت است.')
    add_bullet(doc, 'جلد اطلاعاتی و متعادل با لوگو، عنوان تصمیم و metadata کوتاه')
    add_bullet(doc, 'جدول‌های کم‌تراکم با header تیره، خطوط خاکستری روشن و عرض ستون متناسب')
    add_bullet(doc, 'حذف شماره‌گذاری تزئینی از تیترها و استفاده از hierarchy واقعی Word')


def financial_architecture(doc):
    add_heading(doc, 'معماری اطلاعات مالی پیشنهادی', kicker='بعد')
    add_body(doc, 'قالب عمومی درباره اجزای نرم‌افزار سؤال می‌کند. قالب مالی باید علاوه بر آن، منبع حقیقت و ثابت‌هایی را ثبت کند که خطای آن‌ها به گزارش نادرست ثروت یا افشای داده منجر می‌شود.')

    add_heading(doc, 'شش بخش اختصاصی مالی', level=2)
    sections = [
        ('ثابت‌های مالی', 'تراز هر سند، نبود ستون موجودی، اصلاح با سند معکوس، تفکیک هزینه و انتقال'),
        ('منشأ ارزش', 'منبع نرخ، زمان مؤثر، ارز پایه، snapshot ارزش‌گذاری و سیاست stale data'),
        ('مرز ثبت و برنامه', 'رویداد برنامه‌ریزی‌شده تا زمان اجرا نباید دفترکل را تغییر دهد'),
        ('تطبیق و بازسازی', 'بازسازی مانده از postings، تشخیص اختلاف و شواهد recovery'),
        ('جداسازی کاربر', 'مالکیت ردیف، RLS، tenant context و مسیرهای دسترسی مدیریتی'),
        ('اعتماد عملیاتی', 'RPO، RTO، آزمون restore، کامل بودن قیمت و هشدار شکست کنترل'),
    ]
    add_matrix(doc, ['بخش', 'پرسش اجباری'], sections, [1.7, 5.4], font_size=9)

    add_heading(doc, 'نقشه جریان تصمیم', level=2)
    flow = add_matrix(doc, ['ورودی', 'اعتبارسنجی', 'هسته مالی', 'ذخیره و شواهد'], [
        ('فرم و Server Action', 'کاربر، tenant، واحد و زمان', 'قواعد Decimal، FIFO و تراز', 'Journal Entry، Posting و Audit Log'),
        ('قیمت و نرخ ارز', 'منبع، تاریخ و تازگی', 'ارزش‌گذاری و تبدیل', 'Price Cache و Valuation Snapshot'),
        ('برنامه و تعهد', 'وضعیت و idempotency', 'اجرا در موعد', 'رویداد برنامه و سند اجرای یکتا'),
    ], [1.65, 1.7, 1.75, 2.0], font_size=8.5)

    add_body(doc, 'اصل راهنما  ', bold_lead='اصل راهنما  ')
    add_body(doc, 'هر عددی که در رابط نمایش داده می‌شود باید به رکورد منبع، ارز، زمان مؤثر و مسیر محاسبه قابل ردیابی باشد. در مسیر نوشتن، اثر جانبی فقط پس از ثبت حالت پایدار و بررسی idempotency مجاز است.')


def security_risk(doc):
    add_heading(doc, 'کنترل‌های امنیت و داده', kicker='AUDIT')
    add_body(doc, 'اسناد مخزن استفاده از Supabase Auth، PostgreSQL و RLS را نشان می‌دهند. نسخه مالی قالب باید کنترل‌ها را به شکل قابل اثبات ثبت کند، نه با عبارت‌های کلی مانند امنیت بررسی شود.')
    rows = [
        ('P0', 'جداسازی داده کاربر', 'هر policy به authenticated و شرط مالکیت با auth uid محدود باشد', 'آزمون دو کاربر و query منفی'),
        ('P0', 'کلیدهای privileged', 'secret یا service role هرگز در مرورگر قرار نگیرد', 'بررسی bundle و متغیرهای محیطی'),
        ('P0', 'توابع privileged', 'SECURITY DEFINER فقط در schema غیرعمومی با revoke و کنترل هویت', 'فهرست تابع و grant'),
        ('P1', 'جلسه‌های حساس', 'عملیات بازیابی و مدیریت، session معتبر و rate limit داشته باشد', 'آزمون session منقضی و لغوشده'),
        ('P1', 'Data API', 'exposure و GRANT مستقل از RLS مستند شود', 'فهرست schema و نقش‌ها'),
        ('P1', 'داده مالی در log', 'payload و شناسه‌های حساس از log حذف یا mask شوند', 'نمونه log و تست redaction'),
    ]
    add_matrix(doc, ['اولویت', 'کنترل', 'معیار طراحی', 'مدرک پذیرش'], rows, [0.7, 1.45, 2.85, 2.1], font_size=8.25)

    add_heading(doc, 'ریسک‌های فعلی قالب', level=2)
    risks = [
        ('بالا', 'هویت برند متناقض', 'توازن در رابط و README، وِزان در سند طراحی و تراز در بعضی assetها', 'تعیین canonical brand و rename کامل'),
        ('بالا', 'نبود ثابت‌های مالی', 'امکان تایید طراحی بدون اثبات تراز، idempotency و منشأ ارزش', 'افزودن Financial Invariants به gate'),
        ('متوسط', 'دسترس پذیری سند', 'alt text مفقود و دو پرش سطح تیتر', 'اصلاح ساختار و audit خودکار'),
        ('متوسط', 'صفحه‌های متراکم', 'جدول قرارداد داده و آمادگی عملیاتی در چند صفحه می‌شکنند', 'کاهش placeholder و تقسیم هدفمند'),
    ]
    add_matrix(doc, ['شدت', 'ریسک', 'شاهد', 'اقدام'], risks, [0.8, 1.35, 2.75, 2.2], font_size=8.25)


def rollout(doc):
    add_heading(doc, 'نقشه اجرای بازطراحی', kicker='پیشنهاد اجرایی')
    add_body(doc, 'پیاده‌سازی کامل را می‌توان در چهار گام کوچک و قابل بازگشت انجام داد. هر گام یک خروجی مستقل دارد و فقط پس از عبور از معیار خروج به مرحله بعد می‌رود.')
    rows = [
        ('M1', 'تصمیم هویت', 'نام canonical، لوگو، رنگ و فونت سند', 'تطابق README، UI، manifest و assetها'),
        ('M2', 'بازطراحی قالب', 'جلد، hierarchy، جدول‌ها، footer و نمونه معماری', 'رندر تمیز در Word و PDF و امتیاز a11y بدون High'),
        ('M3', 'افزودن کنترل مالی', 'invariants، provenance، reconciliation، RLS و DR', 'هر کنترل مالک و evidence داشته باشد'),
        ('M4', 'پایلوت RFC واقعی', 'تکمیل قالب با یک تغییر مالی واقعی', 'review کمتر از ۳۰ دقیقه و صفر ابهام P0'),
    ]
    add_matrix(doc, ['مرحله', 'خروجی', 'محتوا', 'معیار خروج'], rows, [0.7, 1.55, 2.65, 2.25], font_size=8.5)

    add_heading(doc, 'معیار پذیرش نسخه حرفه‌ای', level=2)
    gates = [
        'جلد در کمتر از ده ثانیه موضوع، تصمیم، وضعیت و مالک را روشن کند.',
        'هیچ Heading level jump، تصویر بدون alt یا جدول با متن فشرده وجود نداشته باشد.',
        'هر عدد مالی واحد، ارز، زمان مؤثر و منبع داده مشخص داشته باشد.',
        'هر مسیر نوشتن مالی idempotency، رفتار شکست و روش برگشت را ثبت کند.',
        'هر کنترل امنیتی یک مالک و مدرک آزمون یا query قابل تکرار داشته باشد.',
        'صفحه‌های فارسی با RTL واقعی و اعداد جدولی در Word و PDF بدون شکست رندر شوند.',
    ]
    for g in gates:
        add_bullet(doc, g, marker='✓')

    add_heading(doc, 'تصمیم نهایی', level=2)
    add_body(doc, 'بازطراحی توصیه می‌شود. قالب فعلی دور انداخته نشود؛ ستون فقرات فنی آن ارزشمند است. نسخه بعد باید ظاهر محصول را بازتاب دهد و کنترل‌های مالی را از نکته جانبی به معیار پذیرش تبدیل کند. مهم‌ترین پیش‌نیاز، تثبیت نام برند میان توازن، وِزان و تراز است.')


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUT)
    doc = Document(OUT)
    clear_body(doc)
    configure_styles(doc)
    setup_headers_footers(doc)

    cover(doc)
    add_page_break(doc)
    executive_audit(doc)
    add_page_break(doc)
    before_after(doc)
    add_page_break(doc)
    visual_evidence(doc)
    add_page_break(doc)
    financial_architecture(doc)
    add_page_break(doc)
    security_risk(doc)
    add_page_break(doc)
    rollout(doc)

    doc.core_properties.title = 'ممیزی قالب طراحی سیستم برای اپ مالی'
    doc.core_properties.subject = 'مقایسه نسخه موجود و طرح پیشنهادی برای توازن'
    doc.core_properties.author = 'تیم محصول توازن'
    doc.core_properties.keywords = 'System Design, Financial App, Audit, RTL, Supabase, PostgreSQL'
    doc.save(OUT)
    add_alt_text_to_images(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
